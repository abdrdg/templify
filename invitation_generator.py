
# Standard library imports
import sys
import os
import urllib.request
import zipfile
import re
import threading
import json
from datetime import datetime

# Third-party imports
from pdf2image import convert_from_path
from docx2pdf import convert as docx2pdf_convert
from docxtpl import DocxTemplate
import openpyxl
import pandas as pd

# Attendee class for OOP
class Attendee:
	def __init__(self, data_dict):
		self.data = data_dict

	def get_context(self, mapping):
		# mapping: {placeholder: excel_column}
		context = {}
		for ph, col in mapping.items():
			value = self.data.get(col, "")
			# Convert None, NaN, or empty string to empty string and strip whitespace
			if value is None or pd.isna(value) or str(value).strip() == "" or str(value).strip().lower() == "nan":
				context[ph] = ""
			else:
				context[ph] = str(value).strip()
		return context

	def get_filename(self):
		# Use Name or fallback to first column
		name = self.data.get("Name") or list(self.data.values())[0]
		# Clean the name the same way as invitation sender
		cleaned_name = ' '.join(part.replace('.', '').replace('"', "'") for part in str(name).split())
		return cleaned_name


# Modern GUI for invitation generation
import customtkinter as ctk
from tkinter import filedialog, messagebox




# Ensure Poppler is available for pdf2image (Windows only)
def ensure_poppler():
	"""
	Download and extract Poppler for Windows if not already present.
	Returns the bin path containing pdftoppm.exe for pdf2image.
	"""
	if sys.platform != "win32":
		return None
	# Explicitly check the expected path
	poppler_bin = os.path.join(os.path.dirname(__file__), "poppler", "poppler-23.11.0", "Library", "bin")
	pdftoppm_path = os.path.join(poppler_bin, "pdftoppm.exe")
	if os.path.exists(pdftoppm_path):
		return poppler_bin
	# Fallback: search all subfolders for pdftoppm.exe
	poppler_dir = os.path.join(os.path.dirname(__file__), "poppler")
	for root, dirs, files in os.walk(poppler_dir):
		if "pdftoppm.exe" in files:
			return root
	# Download and extract Poppler if not found
	url = "https://github.com/oschwartz10612/poppler-windows/releases/download/v23.11.0-0/Release-23.11.0-0.zip"
	zip_path = os.path.join(poppler_dir, "poppler.zip")
	os.makedirs(poppler_dir, exist_ok=True)
	print("Downloading Poppler...")
	urllib.request.urlretrieve(url, zip_path)
	print("Extracting Poppler...")
	with zipfile.ZipFile(zip_path, 'r') as zip_ref:
		zip_ref.extractall(poppler_dir)
	os.remove(zip_path)
	# Find the extracted folder
	for root, dirs, files in os.walk(poppler_dir):
		if "pdftoppm.exe" in files:
			return root
	return None

class InvitationGeneratorApp(ctk.CTk):

	def __init__(self):
		super().__init__()
		self.title("Invitation Generator")
		self.geometry("1000x700")  # Made wider for two-column layout
		ctk.set_appearance_mode("System")
		ctk.set_default_color_theme("blue")

		# File paths
		self.template_path = ctk.StringVar()
		self.excel_path = ctk.StringVar()
		self.output_folder = ctk.StringVar(value=os.path.abspath("output"))

		# Placeholders and column mapping
		self.placeholders = []
		self.excel_columns = []
		self.mapping_vars = {}
		
		# Fast mode toggle
		self.fast_mode = ctk.BooleanVar(value=False)

		# Initialize generation tracking
		self.tracking_file = "generated_invitations.json"
		self.generated_invitations = self.load_generated_invitations()
		
		# Cancel flag for generation process
		self.is_generating = False
		
		# Pagination for large datasets
		self.items_per_page = 25
		self.current_page = 0
		self.total_pages = 0

		# Initialize selection tracking for invitees
		self.selected_invitees = {}  # Dictionary to track checkbox states
		self.invitees = None  # Will hold the DataFrame of invitees

		# UI Elements
		self.create_widgets()

	def create_widgets(self):
		# Title at the top
		ctk.CTkLabel(self, text="Invitation Generator", font=("Arial", 22)).pack(pady=(10, 0))
		
		# Create two-column layout using a horizontal frame
		columns_frame = ctk.CTkFrame(self, fg_color="transparent")
		columns_frame.pack(fill="both", expand=True, padx=10, pady=(10, 10))
		
		# Left column - Controls
		left_column = ctk.CTkFrame(columns_frame)
		left_column.pack(side="left", fill="both", expand=True, padx=(0, 5))
		
		# Right column - Invitees list
		right_column = ctk.CTkFrame(columns_frame)
		right_column.pack(side="right", fill="both", expand=True, padx=(5, 0))

		# === LEFT COLUMN CONTENT ===
		
		# Template selection
		template_frame = ctk.CTkFrame(left_column)
		template_frame.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(template_frame, text="1. Select DOCX Template:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		template_input_frame = ctk.CTkFrame(template_frame, fg_color="transparent")
		template_input_frame.pack(fill="x", padx=5, pady=(0, 5))
		ctk.CTkEntry(template_input_frame, textvariable=self.template_path, width=300, state="readonly").pack(side="left", padx=(0,5), fill="x", expand=True)
		ctk.CTkButton(template_input_frame, text="Browse", command=self.select_template, width=80).pack(side="right")

		# Excel selection
		excel_frame = ctk.CTkFrame(left_column)
		excel_frame.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(excel_frame, text="2. Select Excel File:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		excel_input_frame = ctk.CTkFrame(excel_frame, fg_color="transparent")
		excel_input_frame.pack(fill="x", padx=5, pady=(0, 5))
		ctk.CTkEntry(excel_input_frame, textvariable=self.excel_path, width=300, state="readonly").pack(side="left", padx=(0,5), fill="x", expand=True)
		ctk.CTkButton(excel_input_frame, text="Browse", command=self.select_excel, width=80).pack(side="right")

		# Mapping area
		mapping_section = ctk.CTkFrame(left_column)
		mapping_section.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(mapping_section, text="3. Map Columns to Placeholders:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		self.mapping_dropdowns_frame = ctk.CTkFrame(mapping_section)
		self.mapping_dropdowns_frame.pack(fill="x", padx=5, pady=(0, 5))

		# Output folder
		output_frame = ctk.CTkFrame(left_column)
		output_frame.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(output_frame, text="4. Output Folder:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		output_input_frame = ctk.CTkFrame(output_frame, fg_color="transparent")
		output_input_frame.pack(fill="x", padx=5, pady=(0, 5))
		ctk.CTkEntry(output_input_frame, textvariable=self.output_folder, width=300, state="readonly").pack(side="left", padx=(0,5), fill="x", expand=True)
		ctk.CTkButton(output_input_frame, text="Change", command=self.select_output_folder, width=80).pack(side="right")

		# Fast mode toggle
		fast_mode_frame = ctk.CTkFrame(left_column)
		fast_mode_frame.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(fast_mode_frame, text="5. Processing Mode:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		fast_toggle_frame = ctk.CTkFrame(fast_mode_frame, fg_color="transparent")
		fast_toggle_frame.pack(fill="x", padx=5, pady=(0, 5))
		self.fast_toggle = ctk.CTkCheckBox(
			fast_toggle_frame, 
			text="Fast Mode (Bulk Processing)", 
			variable=self.fast_mode,
			font=("Arial", 11)
		)
		self.fast_toggle.pack(side="left", padx=5)
		ctk.CTkLabel(
			fast_toggle_frame, 
			text="⚡ Processes all files in batches for better performance", 
			font=("Arial", 9), 
			text_color="gray"
		).pack(side="left", padx=(10, 0))

		# Progress bar
		progress_frame = ctk.CTkFrame(left_column)
		progress_frame.pack(pady=5, fill="x", padx=10)
		ctk.CTkLabel(progress_frame, text="Progress:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		self.progress = ctk.CTkProgressBar(progress_frame)
		self.progress.pack(fill="x", padx=5, pady=(0, 5))
		self.progress.set(0)

		# Status log
		log_frame = ctk.CTkFrame(left_column)
		log_frame.pack(fill="both", expand=True, padx=10, pady=5)
		ctk.CTkLabel(log_frame, text="Status Log:", font=("Arial", 12, "bold")).pack(anchor="w", padx=5)
		self.log_text = ctk.CTkTextbox(log_frame, state="disabled")
		self.log_text.pack(fill="both", expand=True, padx=5, pady=(0, 5))

		# === RIGHT COLUMN CONTENT ===
		
		# Invitees list with status
		ctk.CTkLabel(right_column, text="Invitees Selection", font=("Arial", 16, "bold")).pack(pady=(10, 5))
		
		# Header with refresh and pagination buttons
		header_frame = ctk.CTkFrame(right_column)
		header_frame.pack(fill="x", padx=10, pady=(0, 5))
		
		# Pagination controls
		self.prev_page_btn = ctk.CTkButton(header_frame, text="← Prev", width=80, command=self.prev_page, state="disabled")
		self.prev_page_btn.pack(side="left", padx=5)
		
		self.page_label = ctk.CTkLabel(header_frame, text="Page 1 of 1")
		self.page_label.pack(side="left", padx=10)
		
		self.next_page_btn = ctk.CTkButton(header_frame, text="Next →", width=80, command=self.next_page, state="disabled")
		self.next_page_btn.pack(side="left", padx=5)
		
		self.refresh_btn = ctk.CTkButton(header_frame, text="Refresh", width=80, command=self.update_invitees_list)
		self.refresh_btn.pack(side="right", padx=5)
		
		# Selection buttons frame
		selection_frame = ctk.CTkFrame(right_column)
		selection_frame.pack(fill="x", padx=10, pady=(0, 5))
		self.select_all_btn = ctk.CTkButton(selection_frame, text="Select All", width=80, command=self.select_all_invitees)
		self.select_all_btn.pack(side="left", padx=2)
		self.select_none_btn = ctk.CTkButton(selection_frame, text="Select None", width=80, command=self.select_none_invitees)
		self.select_none_btn.pack(side="left", padx=2)
		self.select_ungenerated_btn = ctk.CTkButton(selection_frame, text="Select New", width=90, command=self.select_ungenerated_invitees)
		self.select_ungenerated_btn.pack(side="left", padx=2)
		
		# Scrollable frame for invitees
		self.invitees_scrollable_frame = ctk.CTkScrollableFrame(right_column)
		self.invitees_scrollable_frame.pack(fill="both", expand=True, padx=10, pady=(0, 5))
		self.invitee_labels = {}  # Store labels for updating
		
		# Generate button at bottom of right column
		generate_frame = ctk.CTkFrame(right_column)
		generate_frame.pack(fill="x", padx=10, pady=(0, 10))
		
		self.generate_btn = ctk.CTkButton(
			generate_frame, 
			text="Generate Invitations", 
			command=self.generate_invitations,
			height=40,
			font=("Arial", 14)
		)
		self.generate_btn.pack(pady=5, fill="x", padx=10)

	def prev_page(self):
		"""Go to previous page"""
		if self.current_page > 0:
			self.current_page -= 1
			self.update_invitees_list()

	def next_page(self):
		"""Go to next page"""
		if self.current_page < self.total_pages - 1:
			self.current_page += 1
			self.update_invitees_list()

	def update_pagination_controls(self):
		"""Update pagination button states and labels"""
		if self.invitees is None or self.invitees.empty:
			self.total_pages = 0
			self.current_page = 0
		else:
			self.total_pages = (len(self.invitees) + self.items_per_page - 1) // self.items_per_page
			if self.current_page >= self.total_pages:
				self.current_page = max(0, self.total_pages - 1)
		
		# Update buttons
		self.prev_page_btn.configure(state="normal" if self.current_page > 0 else "disabled")
		self.next_page_btn.configure(state="normal" if self.current_page < self.total_pages - 1 else "disabled")
		
		# Update label
		if self.total_pages > 0:
			start_item = self.current_page * self.items_per_page + 1
			end_item = min((self.current_page + 1) * self.items_per_page, len(self.invitees))
			self.page_label.configure(text=f"Items {start_item}-{end_item} of {len(self.invitees)} (Page {self.current_page + 1} of {self.total_pages})")
		else:
			self.page_label.configure(text="No items")

	def select_template(self):
		path = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
		if path:
			self.template_path.set(path)
			self.log(f"Selected template: {path}")
			self.placeholders = self.extract_placeholders(path)
			self.update_mapping_dropdowns()

	def extract_placeholders(self, docx_path):
		"""
		Extract placeholders from all XML files in the .docx archive, including those split across runs.
		Returns a list of unique placeholder names.
		"""
		import zipfile
		import xml.etree.ElementTree as ET
		found = set()
		with zipfile.ZipFile(docx_path) as docx_zip:
			for file in docx_zip.namelist():
				if file.endswith('.xml'):
					with docx_zip.open(file) as xml_file:
						try:
							xml = xml_file.read().decode('utf-8')
						except Exception:
							continue
						# Join all <w:t> text nodes for robust placeholder extraction
						try:
							root = ET.fromstring(xml)
							texts = []
							for elem in root.iter():
								# Word text node
								if elem.tag.endswith('}t'):
									texts.append(elem.text or '')
							joined_text = ''.join(texts)
							found.update(re.findall(r'{{\s*(\w+)\s*}}', joined_text))
						except Exception:
							# Fallback: regex on raw xml
							found.update(re.findall(r'{{\s*(\w+)\s*}}', xml))
		self.log(f"Found placeholders: {', '.join(found)}")
		return list(found)

	def select_excel(self):
		path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx;*.xls")])
		if path:
			self.excel_path.set(path)
			self.log(f"Selected Excel: {path}")
			self.excel_columns = self.extract_excel_columns(path)
			self.load_invitees(path)
			# Reset selections when loading new file
			self.reset_all_selections()
			self.current_page = 0  # Reset to first page
			self.update_mapping_dropdowns()
			self.update_invitees_list()

	def load_invitees(self, excel_path):
		"""Load invitees from Excel file"""
		try:
			import pandas as pd
			self.invitees = pd.read_excel(excel_path)
			self.log(f"Loaded {len(self.invitees)} invitees from Excel.")
		except Exception as e:
			self.log(f"Error loading invitees: {e}")
			self.invitees = None

	def extract_excel_columns(self, excel_path):
		wb = openpyxl.load_workbook(excel_path)
		ws = wb.active
		columns = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
		self.log(f"Excel columns: {', '.join([str(c) for c in columns if c])}")
		return [str(c) for c in columns if c]

	def update_mapping_dropdowns(self):
		# Clear previous
		for widget in self.mapping_dropdowns_frame.winfo_children():
			widget.destroy()
		self.mapping_vars = {}
		if not self.placeholders or not self.excel_columns:
			return
		for ph in self.placeholders:
			frame = ctk.CTkFrame(self.mapping_dropdowns_frame)
			frame.pack(fill="x", pady=2)
			ctk.CTkLabel(frame, text=f"{ph}:", width=120).pack(side="left")
			var = ctk.StringVar()
			dropdown = ctk.CTkOptionMenu(frame, variable=var, values=self.excel_columns)
			dropdown.pack(side="left", padx=5)
			self.mapping_vars[ph] = var

	def select_output_folder(self):
		path = filedialog.askdirectory()
		if path:
			self.output_folder.set(path)
			self.log(f"Output folder set to: {path}")

	def update_invitees_list(self):
		"""Update the invitees list with checkboxes - optimized with pagination"""
		self.clear_invitees_list()
		
		if self.invitees is None or self.invitees.empty:
			self.update_pagination_controls()
			return

		# Update pagination controls
		self.update_pagination_controls()
		
		# Calculate which items to show
		start_idx = self.current_page * self.items_per_page
		end_idx = min(start_idx + self.items_per_page, len(self.invitees))
		
		# Only create widgets for visible items
		visible_invitees = self.invitees.iloc[start_idx:end_idx]
		
		# Batch process attendee data first (faster than creating UI widgets)
		attendee_data = []
		for idx, row in visible_invitees.iterrows():
			# Convert row to dict format that Attendee expects
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			
			# Create Attendee object to get consistent filename
			attendee = Attendee(data)
			filename = attendee.get_filename()
			
			# Also get a display name (original, unprocessed)
			name_columns = ['Name', 'name', 'Full Name', 'full_name']
			display_name = None
			for col in name_columns:
				if col in row and pd.notna(row[col]):
					display_name = str(row[col]).strip()
					break
			
			if not display_name:
				# Fallback to first non-null column
				for col_name, value in row.items():
					if pd.notna(value) and str(value).strip():
						display_name = str(value).strip()
						break
				if not display_name:
					display_name = f"Row {idx + 1}"
			
			# Check generation status
			is_generated = self.was_invitation_generated(filename)
			
			attendee_data.append({
				'idx': idx,
				'display_name': display_name,
				'filename': filename,
				'is_generated': is_generated
			})
		
		# Now create UI elements in batch
		for data in attendee_data:
			self._create_invitee_widget(data)

	def _create_invitee_widget(self, data):
		"""Create UI widget for a single invitee"""
		idx = data['idx']
		display_name = data['display_name'] 
		filename = data['filename']
		is_generated = data['is_generated']
		
		# Use the processed filename for the key (consistent with tracking)
		key = f"{idx}|{filename}"
		
		# Create frame for this invitee
		frame = ctk.CTkFrame(self.invitees_scrollable_frame)
		frame.pack(fill="x", padx=2, pady=1)
		
		# Checkbox for selection
		checkbox_var = ctk.BooleanVar()
		checkbox = ctk.CTkCheckBox(frame, text="", variable=checkbox_var, width=20)
		checkbox.pack(side="left", padx=5)
		
		# Store or restore checkbox variable
		if key in self.selected_invitees:
			# Restore previous selection state
			previous_state = self.selected_invitees[key].get()
			checkbox_var.set(previous_state)
		else:
			# Set default selection based on generation status
			if is_generated:
				checkbox_var.set(False)  # Don't select already generated
			else:
				checkbox_var.set(True)  # Select ungenerated by default
		
		# Update/store checkbox variable for later use
		self.selected_invitees[key] = checkbox_var
		
		# Name display (show the original display name)
		info_label = ctk.CTkLabel(frame, text=display_name, anchor="w")
		info_label.pack(side="left", padx=5, fill="x", expand=True)
		
		# Status label
		status_label = ctk.CTkLabel(frame, text="", anchor="e", width=120)
		status_label.pack(side="right", padx=5)
		
		# Store label reference for updates
		self.invitee_labels[key] = status_label
		
		# Update status
		if is_generated:
			status_label.configure(text="Generated ✓", text_color="green")
		else:
			status_label.configure(text="Not generated", text_color="gray")

	def clear_invitees_list(self):
		"""Clear all invitee widgets but preserve selection state"""
		for widget in self.invitees_scrollable_frame.winfo_children():
			widget.destroy()
		self.invitee_labels = {}
		# DON'T clear selected_invitees - preserve selections across pages

	def reset_all_selections(self):
		"""Completely reset all selections (used when loading new Excel file)"""
		self.selected_invitees = {}

	def select_all_invitees(self):
		"""Select all invitees for generation (across all pages)"""
		if self.invitees is None or self.invitees.empty:
			return
			
		count = 0
		# Work with all invitees, not just visible ones
		for idx, row in self.invitees.iterrows():
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			attendee = Attendee(data)
			filename = attendee.get_filename()
			key = f"{idx}|{filename}"
			
			# Create checkbox variable if it doesn't exist
			if key not in self.selected_invitees:
				self.selected_invitees[key] = ctk.BooleanVar()
			
			self.selected_invitees[key].set(True)
			count += 1
			
		self.log(f"Selected all {count} invitees (across all pages).")

	def select_none_invitees(self):
		"""Deselect all invitees (across all pages)"""
		if self.invitees is None or self.invitees.empty:
			return
			
		count = 0
		# Work with all invitees, not just visible ones  
		for idx, row in self.invitees.iterrows():
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			attendee = Attendee(data)
			filename = attendee.get_filename()
			key = f"{idx}|{filename}"
			
			# Create checkbox variable if it doesn't exist
			if key not in self.selected_invitees:
				self.selected_invitees[key] = ctk.BooleanVar()
				
			self.selected_invitees[key].set(False)
			count += 1
			
		self.log(f"Deselected all {count} invitees.")

	def select_ungenerated_invitees(self):
		"""Select only invitees who haven't been generated yet (across all pages)"""
		if self.invitees is None or self.invitees.empty:
			return
			
		selected_count = 0
		total_count = 0
		
		# Work with all invitees, not just visible ones
		for idx, row in self.invitees.iterrows():
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			attendee = Attendee(data)
			filename = attendee.get_filename()
			key = f"{idx}|{filename}"
			
			# Create checkbox variable if it doesn't exist
			if key not in self.selected_invitees:
				self.selected_invitees[key] = ctk.BooleanVar()
			
			total_count += 1
			if not self.was_invitation_generated(filename):
				self.selected_invitees[key].set(True)
				selected_count += 1
			else:
				self.selected_invitees[key].set(False)
				
		self.log(f"Selected {selected_count} ungenerated invitees out of {total_count} total.")

	def log(self, message):
		# Schedule the log update on the main thread
		self.after(0, self._log_update, message)
	
	def _log_update(self, message):
		self.log_text.configure(state="normal")
		self.log_text.insert("end", message + "\n")
		self.log_text.see("end")
		self.log_text.configure(state="disabled")

	def load_generated_invitations(self):
		"""Load the record of generated invitations from JSON file"""
		if os.path.exists(self.tracking_file):
			try:
				with open(self.tracking_file, 'r') as f:
					return json.load(f)
			except json.JSONDecodeError:
				self.log("Warning: Generation tracking file corrupted, starting fresh.")
				return {}
		return {}

	def save_generated_invitations(self):
		"""Save the record of generated invitations to JSON file"""
		try:
			with open(self.tracking_file, 'w') as f:
				json.dump(self.generated_invitations, f, indent=2)
		except Exception as e:
			self.log(f"Warning: Could not save generation tracking: {e}")

	def was_invitation_generated(self, name):
		"""Check if invitation was already generated for this person"""
		return name in self.generated_invitations

	def find_existing_invitation_files(self, name):
		"""Find existing invitation files, trying different filename variations for backward compatibility"""
		output_folder = self.output_folder.get()
		if not output_folder or not os.path.exists(output_folder):
			return None
			
		# Try the current cleaned filename first
		cleaned_name = self.get_filename_from_name(name)
		primary_files = {
			'docx': os.path.join(output_folder, f"Invitation - {cleaned_name}.docx"),
			'pdf': os.path.join(output_folder, f"Invitation - {cleaned_name}.pdf"),
			'png': os.path.join(output_folder, f"Invitation - {cleaned_name}.png")
		}
		
		if os.path.exists(primary_files['docx']):
			return primary_files
		
		# Try various legacy processing variations for backward compatibility
		variations = [
			# Legacy processing (old get_filename logic)
			str(name).replace("\n", " ").replace(".", "").replace('"', "'").strip(),
			# Raw name with just quote replacement
			str(name).replace('"', "'"),
			# Raw name with no processing
			str(name),
			# Name with just space normalization (but keeping leading/trailing)
			str(name).replace("\n", " ").replace(".", "").replace('"', "'"),
		]
		
		for variation in variations:
			# Try exact variation
			files = {
				'docx': os.path.join(output_folder, f"Invitation - {variation}.docx"),
				'pdf': os.path.join(output_folder, f"Invitation - {variation}.pdf"),
				'png': os.path.join(output_folder, f"Invitation - {variation}.png")
			}
			if os.path.exists(files['docx']):
				return files
				
			# Also try with potential extra space after dash (leading space in name)
			files_extra_space = {
				'docx': os.path.join(output_folder, f"Invitation -  {variation}.docx"),
				'pdf': os.path.join(output_folder, f"Invitation -  {variation}.pdf"),
				'png': os.path.join(output_folder, f"Invitation -  {variation}.png")
			}
			if os.path.exists(files_extra_space['docx']):
				return files_extra_space
		
		# Try fuzzy matching as last resort
		name_parts = cleaned_name.lower().split()
		if name_parts and os.path.exists(output_folder):
			try:
				for filename in os.listdir(output_folder):
					if filename.startswith("Invitation - ") and filename.endswith(".docx"):
						file_name_part = filename[13:-5].lower()  # Remove "Invitation - " and ".docx"
						# Check if all name parts are present in the filename
						if all(part in file_name_part for part in name_parts):
							base_path = os.path.join(output_folder, filename[:-5])  # Remove .docx
							fuzzy_files = {
								'docx': base_path + '.docx',
								'pdf': base_path + '.pdf',
								'png': base_path + '.png'
							}
							return fuzzy_files
			except OSError:
				pass  # Handle permission errors gracefully
		
		return None

	def get_filename_from_name(self, name):
		"""Get filename from name using the same logic as Attendee.get_filename()"""
		return ' '.join(part.replace('.', '').replace('"', "'") for part in str(name).replace('\n', ' ').split())

	def mark_invitation_generated(self, name, output_folder):
		"""Mark invitation as generated for this person"""
		self.generated_invitations[name] = {
			"generated_date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
			"output_folder": output_folder
		}
		self.save_generated_invitations()

	def generate_invitations(self):
		if self.is_generating:
			# Cancel generation
			self.is_generating = False
			self.log("Generation cancelled by user.")
			self.reset_generate_button()
			return
			
		# Start generation in a separate thread to keep UI responsive
		self.is_generating = True
		self.generate_btn.configure(text="Cancel", fg_color="red")
		thread = threading.Thread(target=self._generate_invitations_thread, daemon=True)
		thread.start()

	def reset_generate_button(self):
		"""Reset the generate button to its original state"""
		self.is_generating = False
		self.generate_btn.configure(text="Generate Invitations", fg_color=["#1f538d", "#14375e"])

	def _generate_invitations_thread(self):
		try:
			self._do_generation()
		finally:
			# Always reset the button when generation ends
			self.after(0, self.reset_generate_button)
	
	def _do_generation(self):
		self.log("Generation started...")
		template_path = self.template_path.get()
		excel_path = self.excel_path.get()
		output_folder = self.output_folder.get()
		if not template_path or not excel_path or not output_folder:
			self.log("Please select a template, Excel file, and output folder.")
			return

		# Build mapping from dropdowns
		mapping = {ph: var.get() for ph, var in self.mapping_vars.items()}
		if not all(mapping.values()):
			self.log("Please map all placeholders to Excel columns.")
			return

		# Check if we have invitees loaded and selected
		if self.invitees is None or self.invitees.empty:
			self.log("No invitees data loaded.")
			return

		# Count selected invitees
		selected_count = 0
		selected_indices = []
		for key, checkbox_var in self.selected_invitees.items():
			if checkbox_var.get():
				idx, name = key.split("|", 1)
				selected_indices.append(int(idx))
				selected_count += 1

		if selected_count == 0:
			self.log("No invitees selected for generation.")
			return

		self.log(f"Starting to generate {selected_count} selected invitations...")
		
		# Check if fast mode is enabled
		if self.fast_mode.get():
			self._generate_fast_mode(template_path, output_folder, mapping, selected_indices, selected_count)
		else:
			self._generate_normal_mode(template_path, output_folder, mapping, selected_indices, selected_count)

	def _generate_fast_mode(self, template_path, output_folder, mapping, selected_indices, selected_count):
		"""Fast mode: Process in bulk stages - DOCX, then PDF, then PNG"""
		self.log("🚀 Fast mode enabled - Processing in bulk stages...")
		
		# Prepare output folder
		os.makedirs(output_folder, exist_ok=True)
		
		# Ensure Poppler is available for pdf2image
		poppler_path = None
		if sys.platform == "win32":
			poppler_path = ensure_poppler()
		
		generated_files = []
		docx_files = []
		pdf_files = []
		
		# STAGE 1: Generate all DOCX files
		self.log("📄 Stage 1/3: Generating DOCX files...")
		docx_generated = 0
		
		for i, idx in enumerate(selected_indices):
			if not self.is_generating:
				self.log("Generation cancelled.")
				return
				
			row = self.invitees.iloc[idx]
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			
			attendee = Attendee(data)
			context = attendee.get_context(mapping)
			filename = attendee.get_filename()
			
			try:
				doc = DocxTemplate(template_path)
				doc.render(context)
				out_docx = os.path.join(output_folder, f"Invitation - {filename}.docx")
				doc.save(out_docx)
				
				docx_files.append(out_docx)
				generated_files.append(filename)
				docx_generated += 1
				
				# Update progress
				progress = (i + 1) / (selected_count * 3)  # 3 stages total
				self.after(0, self.progress.set, progress)
				
			except Exception as e:
				self.log(f"Error creating DOCX for {filename}: {e}")
		
		self.log(f"✅ Stage 1 complete: {docx_generated}/{selected_count} DOCX files created")
		
		# STAGE 2: Convert all DOCX to PDF
		if docx_files:
			self.log("📑 Stage 2/3: Converting DOCX to PDF...")
			pdf_converted = 0
			
			try:
				# Use batch processing - much more efficient!
				# docx2pdf can convert an entire directory at once
				self.log("Using batch conversion for better performance...")
				docx2pdf_convert(output_folder, output_folder)
				
				# Check which PDFs were actually created
				for docx_path in docx_files:
					pdf_path = docx_path.replace('.docx', '.pdf')
					if os.path.exists(pdf_path):
						pdf_files.append(pdf_path)
						pdf_converted += 1
				
				# Update progress for the entire batch
				progress = (selected_count * 2) / (selected_count * 3)
				self.after(0, self.progress.set, progress)
				
			except Exception as e:
				self.log(f"Batch PDF conversion failed, falling back to individual conversion: {e}")
				
				# Fallback to individual file conversion
				for i, docx_path in enumerate(docx_files):
					if not self.is_generating:
						self.log("Generation cancelled.")
						return
						
					try:
						# Get the corresponding PDF path
						pdf_path = docx_path.replace('.docx', '.pdf')
						docx2pdf_convert(docx_path, output_folder)
						
						if os.path.exists(pdf_path):
							pdf_files.append(pdf_path)
							pdf_converted += 1
						
						# Update progress
						progress = (selected_count + i + 1) / (selected_count * 3)
						self.after(0, self.progress.set, progress)
						
					except Exception as e:
						self.log(f"Error converting to PDF: {os.path.basename(docx_path)} - {e}")
			
			self.log(f"✅ Stage 2 complete: {pdf_converted}/{len(docx_files)} PDF files created")
		
		# STAGE 3: Convert all PDF to PNG
		if pdf_files:
			self.log("🖼️ Stage 3/3: Converting PDF to PNG...")
			png_converted = 0
			
			for i, pdf_path in enumerate(pdf_files):
				if not self.is_generating:
					self.log("Generation cancelled.")
					return
					
				try:
					png_path = pdf_path.replace('.pdf', '.png')
					images = convert_from_path(pdf_path, dpi=200, fmt='png', poppler_path=poppler_path)
					
					if images:
						images[0].save(png_path, 'PNG')
						png_converted += 1
					
					# Update progress
					progress = (selected_count * 2 + i + 1) / (selected_count * 3)
					self.after(0, self.progress.set, progress)
					
				except Exception as e:
					self.log(f"Error converting to PNG: {os.path.basename(pdf_path)} - {e}")
			
			self.log(f"✅ Stage 3 complete: {png_converted}/{len(pdf_files)} PNG files created")
		
		# Mark all generated files and update UI
		for filename in generated_files:
			self.mark_invitation_generated(filename, output_folder)
		
		# Update invitee statuses
		for idx in selected_indices:
			row = self.invitees.iloc[idx]
			data = row.to_dict()
			data = {str(k): v for k, v in data.items()}
			attendee = Attendee(data)
			filename = attendee.get_filename()
			key = f"{idx}|{filename}"
			if key in self.invitee_labels:
				self.after(0, self.update_invitee_status, key, True)
		
		self.log(f"🎉 Fast mode generation complete! Generated: {len(generated_files)} invitations")
		self.after(0, self.update_invitees_list)

	def _generate_normal_mode(self, template_path, output_folder, mapping, selected_indices, selected_count):
		"""Normal mode: Process each invitation completely before moving to the next"""
		self.log("🐌 Normal mode: Processing each invitation completely...")
		
		# Prepare output folder
		os.makedirs(output_folder, exist_ok=True)

		# Ensure Poppler is available for pdf2image
		poppler_path = None
		if sys.platform == "win32":
			poppler_path = ensure_poppler()

		# Generate invitations for selected invitees only
		generated_count = 0
		current_processed = 0

		for idx in selected_indices:
			# Check for cancellation
			if not self.is_generating:
				self.log("Generation cancelled.")
				return
				
			current_processed += 1
			row = self.invitees.iloc[idx]
			data = row.to_dict()
			
			# Convert data keys to strings for consistency
			data = {str(k): v for k, v in data.items()}
			
			attendee = Attendee(data)
			context = attendee.get_context(mapping)
			filename = attendee.get_filename()
			
			try:
				doc = DocxTemplate(template_path)
				doc.render(context)
				out_docx = os.path.join(output_folder, f"Invitation - {filename}.docx")
				out_pdf = os.path.join(output_folder, f"Invitation - {filename}.pdf")
				out_png = os.path.join(output_folder, f"Invitation - {filename}.png")
				doc.save(out_docx)
				self.log(f"Saved: {out_docx}")
				
				# Convert DOCX to PDF
				pdf_success = False
				try:
					docx2pdf_convert(out_docx, output_folder)
					self.log(f"PDF created: {out_pdf}")
					pdf_success = True
				except Exception as e:
					self.log(f"PDF conversion failed: {e}")
					out_pdf = None
				
				# Convert PDF to PNG (first page)
				png_success = False
				if out_pdf and os.path.exists(out_pdf):
					try:
						images = convert_from_path(out_pdf, dpi=200, fmt='png', poppler_path=poppler_path)
						if images:
							images[0].save(out_png, 'PNG')
							self.log(f"PNG created: {out_png}")
							png_success = True
					except Exception as e:
						self.log(f"PNG conversion failed: {e}")
				
				# Mark as generated only if at least the DOCX was created successfully
				self.mark_invitation_generated(filename, output_folder)
				generated_count += 1
				
				# Store status update for later batch processing
				key = f"{idx}|{filename}"
				if key in self.invitee_labels:
					self.after(0, self.update_invitee_status, key, True)
				
			except Exception as e:
				self.log(f"Error for {filename}: {e}")
			
			# Update progress and log every 10 items to reduce UI updates
			if current_processed % 10 == 0 or current_processed == selected_count:
				self.after(0, self.progress.set, current_processed / selected_count)
				if current_processed % 10 == 0:
					self.log(f"Progress: {current_processed}/{selected_count} processed")

		self.log(f"Generation complete. Generated: {generated_count} invitations")
		
		# Only refresh the current page to show updated statuses
		self.after(0, self.update_invitees_list)

	def update_invitee_status(self, key, is_generated):
		"""Update the status display for a single invitee"""
		if key in self.invitee_labels:
			label = self.invitee_labels[key]
			if is_generated:
				label.configure(text="Generated", text_color="green")
			else:
				label.configure(text="Not generated", text_color="gray")

if __name__ == "__main__":
	app = InvitationGeneratorApp()
	app.mainloop()
