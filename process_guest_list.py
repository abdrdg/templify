import pandas as pd
import json
import random
import string
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import customtkinter as ctk

class GuestListProcessor:
    def __init__(self):
        self.guest_data_file = "guest_codes.json"
        self.guest_tracking_file = "guest_tracking.json"
        self.guest_data = self.load_existing_data()
        self.tracking_data = self.load_tracking_data()
        
    def load_existing_data(self):
        """Load existing guest code assignments from JSON file"""
        if os.path.exists(self.guest_data_file):
            try:
                with open(self.guest_data_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except (json.JSONDecodeError, FileNotFoundError):
                return {}
        return {}
    
    def load_tracking_data(self):
        """Load tracking data for email sending status"""
        if os.path.exists(self.guest_tracking_file):
            try:
                with open(self.guest_tracking_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except (json.JSONDecodeError, FileNotFoundError):
                return {}
        return {}
    
    def save_data(self):
        """Save guest data to JSON file"""
        try:
            with open(self.guest_data_file, 'w', encoding='utf-8') as f:
                json.dump(self.guest_data, f, indent=2, ensure_ascii=False)
        except Exception as e:
            print(f"Error saving guest data: {e}")
    
    def save_tracking_data(self):
        """Save tracking data to JSON file"""
        try:
            with open(self.guest_tracking_file, 'w', encoding='utf-8') as f:
                json.dump(self.tracking_data, f, indent=2, ensure_ascii=False)
        except Exception as e:
            print(f"Error saving tracking data: {e}")
    
    def generate_unique_code(self):
        """Generate a unique 6-digit alphanumeric code"""
        while True:
            # Generate 6-digit alphanumeric code (uppercase letters and numbers)
            code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
            
            # Check if code already exists
            if not any(guest.get('unique_code') == code for guest in self.guest_data.values()):
                return code
    
    def create_guest_key(self, name, email):
        """Create a unique key for each guest based on name and email"""
        return f"{str(name).strip().lower()}_{str(email).strip().lower()}"
    
    def process_excel_file(self, excel_path, name_column='Name', email_column='Email'):
        """
        Process Excel file and generate unique codes for new guests
        
        Args:
            excel_path: Path to Excel file
            name_column: Column name containing guest names
            email_column: Column name containing guest emails
        
        Returns:
            dict: Processing results with counts
        """
        try:
            # Read Excel file
            df = pd.read_excel(excel_path)
            
            # Check if required columns exist
            if name_column not in df.columns:
                raise ValueError(f"Column '{name_column}' not found in Excel file")
            if email_column not in df.columns:
                raise ValueError(f"Column '{email_column}' not found in Excel file")
            
            results = {
                'total_guests': len(df),
                'new_codes_generated': 0,
                'existing_codes': 0,
                'invalid_entries': 0,
                'processed_guests': []
            }
            
            for index, row in df.iterrows():
                name = str(row[name_column]).strip() if pd.notna(row[name_column]) else ''
                email = str(row[email_column]).strip() if pd.notna(row[email_column]) else ''
                
                # Skip if name or email is empty
                if not name or not email or name.lower() == 'nan' or email.lower() == 'nan':
                    results['invalid_entries'] += 1
                    continue
                
                # Create unique key for this guest
                guest_key = self.create_guest_key(name, email)
                
                # Check if guest already has a code
                if guest_key in self.guest_data:
                    results['existing_codes'] += 1
                    guest_info = self.guest_data[guest_key].copy()
                else:
                    # Generate new code for new guest
                    unique_code = self.generate_unique_code()
                    guest_info = {
                        'name': name,
                        'email': email,
                        'unique_code': unique_code,
                        'date_generated': datetime.now().isoformat(),
                        'email_sent': False,
                        'date_email_sent': None
                    }
                    
                    # Save to guest data
                    self.guest_data[guest_key] = guest_info
                    results['new_codes_generated'] += 1
                
                # Add any additional columns from Excel
                guest_info['excel_row_data'] = {}
                for col in df.columns:
                    if col not in [name_column, email_column]:
                        guest_info['excel_row_data'][col] = str(row[col]) if pd.notna(row[col]) else ''
                
                results['processed_guests'].append(guest_info)
            
            # Save data
            self.save_data()
            
            return results
            
        except Exception as e:
            raise Exception(f"Error processing Excel file: {str(e)}")
    
    def export_to_excel(self, output_path="guest_list_with_codes.xlsx"):
        """Export guest data to Excel file for staff use during event"""
        try:
            # Create DataFrame from guest data
            guests_list = []
            for guest_key, guest_info in self.guest_data.items():
                guest_row = {
                    'Name': guest_info['name'],
                    'Email': guest_info['email'],
                    'Unique_Code': guest_info['unique_code'],
                    'Date_Generated': guest_info['date_generated'],
                    'Email_Sent': guest_info['email_sent'],
                    'Date_Email_Sent': guest_info.get('date_email_sent', ''),
                    'Checked_In': '',  # For staff to mark during event
                    'Check_In_Time': '',  # For staff to record check-in time
                }
                
                # Add any additional Excel data
                if 'excel_row_data' in guest_info:
                    guest_row.update(guest_info['excel_row_data'])
                
                guests_list.append(guest_row)
            
            # Sort by name
            guests_list.sort(key=lambda x: x['Name'])
            
            # Create DataFrame
            df = pd.DataFrame(guests_list)
            
            # Create Excel file with formatting
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Guest List', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Guest List']
                
                # Define styles
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                center_alignment = Alignment(horizontal="center", vertical="center")
                
                # Apply header formatting
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_alignment
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add borders and alternating row colors
                from openpyxl.styles import Border, Side
                thin_border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                light_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                
                for row_num, row in enumerate(worksheet.iter_rows(min_row=1, max_row=len(guests_list) + 1), 1):
                    for cell in row:
                        cell.border = thin_border
                        if row_num > 1 and row_num % 2 == 0:  # Alternating rows (skip header)
                            cell.fill = light_fill
            
            return output_path
            
        except Exception as e:
            raise Exception(f"Error exporting to Excel: {str(e)}")
    
    def export_to_json(self, output_path="guest_data.json"):
        """Export guest data to JSON file"""
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(self.guest_data, f, indent=2, ensure_ascii=False)
            return output_path
        except Exception as e:
            raise Exception(f"Error exporting to JSON: {str(e)}")
    
    def send_codes_by_email(self, smtp_server, smtp_port, email_user, email_password, 
                           subject_template="Your Event Access Code", 
                           body_template=None, progress_callback=None):
        """
        Send unique codes to guests via email
        
        Args:
            smtp_server: SMTP server address
            smtp_port: SMTP server port
            email_user: Sender email
            email_password: Sender email password
            subject_template: Email subject template
            body_template: Email body template (can include {name} and {code} placeholders)
            progress_callback: Function to call with progress updates
        """
        if body_template is None:
            body_template = """
Dear {name},

Your unique access code for the event is: {code}

Please present this code at the event check-in.

Best regards,
Event Organizers
"""
        
        sent_count = 0
        failed_count = 0
        
        try:
            # Connect to SMTP server
            server = smtplib.SMTP(smtp_server, smtp_port)
            server.starttls()
            server.login(email_user, email_password)
            
            total_to_send = sum(1 for guest in self.guest_data.values() if not guest.get('email_sent', False))
            
            for guest_key, guest_info in self.guest_data.items():
                if guest_info.get('email_sent', False):
                    continue  # Skip already sent emails
                
                try:
                    # Create email message
                    msg = MIMEMultipart()
                    msg['From'] = email_user
                    msg['To'] = guest_info['email']
                    msg['Subject'] = subject_template
                    
                    # Format email body
                    body = body_template.format(
                        name=guest_info['name'],
                        code=guest_info['unique_code']
                    )
                    
                    msg.attach(MIMEText(body, 'plain'))
                    
                    # Send email
                    server.send_message(msg)
                    
                    # Update guest info
                    guest_info['email_sent'] = True
                    guest_info['date_email_sent'] = datetime.now().isoformat()
                    
                    sent_count += 1
                    
                    if progress_callback:
                        progress_callback(sent_count, total_to_send, guest_info['name'])
                
                except Exception as e:
                    failed_count += 1
                    print(f"Failed to send email to {guest_info['name']} ({guest_info['email']}): {e}")
            
            server.quit()
            
            # Save updated data
            self.save_data()
            
            return {
                'sent': sent_count,
                'failed': failed_count,
                'total_processed': sent_count + failed_count
            }
            
        except Exception as e:
            raise Exception(f"Email sending error: {str(e)}")
    
    def get_stats(self):
        """Get statistics about the guest list"""
        total_guests = len(self.guest_data)
        emails_sent = sum(1 for guest in self.guest_data.values() if guest.get('email_sent', False))
        emails_pending = total_guests - emails_sent
        
        return {
            'total_guests': total_guests,
            'emails_sent': emails_sent,
            'emails_pending': emails_pending,
            'unique_codes_generated': total_guests
        }
    
    def verify_code(self, code):
        """Verify if a code exists and return guest information"""
        for guest_key, guest_info in self.guest_data.items():
            if guest_info.get('unique_code') == code.upper():
                return guest_info
        return None
    
    def get_guest_by_code(self, code):
        """Get guest information by unique code"""
        return self.verify_code(code)
    
    def mark_checked_in(self, code):
        """Mark a guest as checked in"""
        guest = self.verify_code(code)
        if guest:
            # Update in memory
            for guest_key, guest_info in self.guest_data.items():
                if guest_info.get('unique_code') == code.upper():
                    guest_info['checked_in'] = True
                    guest_info['check_in_time'] = datetime.now().isoformat()
                    break
            
            # Save data
            self.save_data()
            return True
        return False


# GUI Application using CustomTkinter
class GuestListProcessorGUI(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        self.processor = GuestListProcessor()
        
        self.title("Guest List Processor")
        self.geometry("1000x700")
        self.minsize(800, 600)
        
        # Configure grid
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        
        self.create_widgets()
        self.update_stats()
    
    def create_widgets(self):
        # Title
        title_label = ctk.CTkLabel(self, text="Guest List Processor", font=ctk.CTkFont(size=24, weight="bold"))
        title_label.grid(row=0, column=0, padx=20, pady=20)
        
        # Main frame
        main_frame = ctk.CTkFrame(self)
        main_frame.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="nsew")
        main_frame.grid_columnconfigure((0, 1), weight=1)
        main_frame.grid_rowconfigure(1, weight=1)
        
        # Left column - File operations
        left_frame = ctk.CTkFrame(main_frame)
        left_frame.grid(row=0, column=0, rowspan=2, padx=10, pady=10, sticky="nsew")
        
        # File selection
        file_label = ctk.CTkLabel(left_frame, text="Excel File Operations", font=ctk.CTkFont(size=16, weight="bold"))
        file_label.pack(padx=20, pady=(20, 10))
        
        self.file_path_var = tk.StringVar()
        file_entry = ctk.CTkEntry(left_frame, textvariable=self.file_path_var, width=300)
        file_entry.pack(padx=20, pady=5)
        
        browse_btn = ctk.CTkButton(left_frame, text="Browse Excel File", command=self.browse_file)
        browse_btn.pack(padx=20, pady=5)
        
        # Column selection
        col_frame = ctk.CTkFrame(left_frame)
        col_frame.pack(padx=20, pady=10, fill="x")
        
        ctk.CTkLabel(col_frame, text="Name Column:").pack(padx=10, pady=5)
        self.name_col_var = ctk.StringVar(value="Name")
        name_entry = ctk.CTkEntry(col_frame, textvariable=self.name_col_var)
        name_entry.pack(padx=10, pady=5)
        
        ctk.CTkLabel(col_frame, text="Email Column:").pack(padx=10, pady=5)
        self.email_col_var = ctk.StringVar(value="Email")
        email_entry = ctk.CTkEntry(col_frame, textvariable=self.email_col_var)
        email_entry.pack(padx=10, pady=5)
        
        process_btn = ctk.CTkButton(left_frame, text="Process Guest List", command=self.process_file)
        process_btn.pack(padx=20, pady=10)
        
        # Export buttons
        export_frame = ctk.CTkFrame(left_frame)
        export_frame.pack(padx=20, pady=10, fill="x")
        
        ctk.CTkLabel(export_frame, text="Export Options", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=5)
        
        export_excel_btn = ctk.CTkButton(export_frame, text="Export to Excel", command=self.export_excel)
        export_excel_btn.pack(padx=10, pady=5)
        
        export_json_btn = ctk.CTkButton(export_frame, text="Export to JSON", command=self.export_json)
        export_json_btn.pack(padx=10, pady=5)
        
        # Right column - Statistics and email
        right_frame = ctk.CTkFrame(main_frame)
        right_frame.grid(row=0, column=1, padx=10, pady=10, sticky="ew")
        
        stats_label = ctk.CTkLabel(right_frame, text="Statistics", font=ctk.CTkFont(size=16, weight="bold"))
        stats_label.pack(padx=20, pady=(20, 10))
        
        self.stats_frame = ctk.CTkFrame(right_frame)
        self.stats_frame.pack(padx=20, pady=5, fill="x")
        
        # Email settings frame
        email_frame = ctk.CTkFrame(main_frame)
        email_frame.grid(row=1, column=1, padx=10, pady=10, sticky="nsew")
        
        email_label = ctk.CTkLabel(email_frame, text="Email Settings", font=ctk.CTkFont(size=16, weight="bold"))
        email_label.pack(padx=20, pady=(20, 10))
        
        # SMTP settings
        smtp_frame = ctk.CTkFrame(email_frame)
        smtp_frame.pack(padx=20, pady=5, fill="x")
        
        # Simple email settings for demo
        ctk.CTkLabel(smtp_frame, text="SMTP Server:").pack(padx=5, pady=2)
        self.smtp_server_var = ctk.StringVar(value="smtp.gmail.com")
        ctk.CTkEntry(smtp_frame, textvariable=self.smtp_server_var, width=200).pack(padx=5, pady=2)
        
        ctk.CTkLabel(smtp_frame, text="SMTP Port:").pack(padx=5, pady=2)
        self.smtp_port_var = ctk.StringVar(value="587")
        ctk.CTkEntry(smtp_frame, textvariable=self.smtp_port_var, width=200).pack(padx=5, pady=2)
        
        ctk.CTkLabel(smtp_frame, text="Email:").pack(padx=5, pady=2)
        self.email_var = ctk.StringVar()
        ctk.CTkEntry(smtp_frame, textvariable=self.email_var, width=200).pack(padx=5, pady=2)
        
        ctk.CTkLabel(smtp_frame, text="Password:").pack(padx=5, pady=2)
        self.password_var = ctk.StringVar()
        password_entry = ctk.CTkEntry(smtp_frame, textvariable=self.password_var, show="*", width=200)
        password_entry.pack(padx=5, pady=2)
        
        send_emails_btn = ctk.CTkButton(email_frame, text="Send Codes via Email", command=self.send_emails)
        send_emails_btn.pack(padx=20, pady=10)
        
        # Progress bar
        self.progress_var = tk.StringVar()
        self.progress_label = ctk.CTkLabel(email_frame, textvariable=self.progress_var)
        self.progress_label.pack(padx=20, pady=5)
        
    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        if filename:
            self.file_path_var.set(filename)
    
    def process_file(self):
        if not self.file_path_var.get():
            messagebox.showerror("Error", "Please select an Excel file")
            return
        
        try:
            results = self.processor.process_excel_file(
                self.file_path_var.get(),
                self.name_col_var.get(),
                self.email_col_var.get()
            )
            
            message = f"""Processing Complete!
            
Total Guests: {results['total_guests']}
New Codes Generated: {results['new_codes_generated']}
Existing Codes: {results['existing_codes']}
Invalid Entries: {results['invalid_entries']}"""
            
            messagebox.showinfo("Success", message)
            self.update_stats()
            
        except Exception as e:
            messagebox.showerror("Error", f"Error processing file: {str(e)}")
    
    def export_excel(self):
        try:
            filename = filedialog.asksaveasfilename(
                title="Save Excel File",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
            )
            if filename:
                self.processor.export_to_excel(filename)
                messagebox.showinfo("Success", f"Excel file exported successfully to:\n{filename}")
        except Exception as e:
            messagebox.showerror("Error", f"Error exporting Excel: {str(e)}")
    
    def export_json(self):
        try:
            filename = filedialog.asksaveasfilename(
                title="Save JSON File",
                defaultextension=".json",
                filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
            )
            if filename:
                self.processor.export_to_json(filename)
                messagebox.showinfo("Success", f"JSON file exported successfully to:\n{filename}")
        except Exception as e:
            messagebox.showerror("Error", f"Error exporting JSON: {str(e)}")
    
    def send_emails(self):
        if not all([self.smtp_server_var.get(), self.email_var.get(), self.password_var.get()]):
            messagebox.showerror("Error", "Please fill in all email settings")
            return
        
        def progress_callback(sent, total, guest_name):
            self.progress_var.set(f"Sending... {sent}/{total} - {guest_name}")
            self.update()
        
        def send_in_thread():
            try:
                results = self.processor.send_codes_by_email(
                    self.smtp_server_var.get(),
                    int(self.smtp_port_var.get()),
                    self.email_var.get(),
                    self.password_var.get(),
                    progress_callback=progress_callback
                )
                
                self.after(0, lambda: self.email_complete(results))
                
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Email sending failed: {str(e)}"))
                self.after(0, lambda: self.progress_var.set(""))
        
        # Start sending in background thread
        threading.Thread(target=send_in_thread, daemon=True).start()
    
    def email_complete(self, results):
        message = f"""Email Sending Complete!
        
Emails Sent: {results['sent']}
Failed: {results['failed']}
Total Processed: {results['total_processed']}"""
        
        messagebox.showinfo("Email Results", message)
        self.progress_var.set("")
        self.update_stats()
    
    def update_stats(self):
        stats = self.processor.get_stats()
        
        # Clear existing stats
        for widget in self.stats_frame.winfo_children():
            widget.destroy()
        
        # Add current stats
        ctk.CTkLabel(self.stats_frame, text=f"Total Guests: {stats['total_guests']}").pack(padx=10, pady=2)
        ctk.CTkLabel(self.stats_frame, text=f"Codes Generated: {stats['unique_codes_generated']}").pack(padx=10, pady=2)
        ctk.CTkLabel(self.stats_frame, text=f"Emails Sent: {stats['emails_sent']}").pack(padx=10, pady=2)
        ctk.CTkLabel(self.stats_frame, text=f"Emails Pending: {stats['emails_pending']}").pack(padx=10, pady=2)


def main():
    """Main function to run the GUI application"""
    ctk.set_appearance_mode("System")
    ctk.set_default_color_theme("blue")
    
    app = GuestListProcessorGUI()
    app.mainloop()


if __name__ == "__main__":
    main()